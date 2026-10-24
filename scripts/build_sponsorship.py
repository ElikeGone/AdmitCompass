"""Filter the DOL OFLC H-1B LCA disclosure file to a small, auditable employer-sponsorship dataset.

Input (not committed; see data/sponsorship/README.md):  data/sponsorship/raw/LCA_Disclosure_Data_FY2025_Q1..Q4.xlsx
  DOL publishes one file per quarter (each covers that quarter's decisions only — verified: the Q4 file's DECISION_DATE
  values are all Jul–Sep 2025). All quarter files present are combined; provenance records exactly which were used.
Outputs:
  data/sponsorship/lca_filtered.csv     one row per employer + occupation + metro + fiscal year (filings, positions, median wage)
  data/sponsorship/provenance.json      source URL, file name, sha256, fiscal year, download date, filter rules, row counts

Column names were read from the file's header row before this script was written (98 columns; see provenance.json).

STATED ASSUMPTIONS (documented, not facts about any person or employer):
  Visa class: H-1B only (the file also contains E-3 Australian and H-1B1 Chile/Singapore).
  Status:     CASE_STATUS == "Certified" only ("Certified - Withdrawn", "Withdrawn" and "Denied" are excluded).
  Occupations (there is no SOC code for product management, so it is defined by job title):
    management_consulting      SOC 13-1111 Management Analysts
    corporate_finance_strategy SOC 13-2051 Financial and Investment Analysts, 13-2054 Financial Risk Specialists,
                               11-3031 Financial Managers; plus SOC 13-1199 Business Operations Specialists, All Other
                               when JOB_TITLE contains "strateg"
    product_management         JOB_TITLE matches "product manager/management/lead/owner/director" (not "product marketing"),
                               any SOC in major groups 11 (management), 13 (business/financial) or 15 (computer/math)
    A case matching two definitions is counted under the first one in the order above.
  Metros (worksite city + state, upper-cased exact match):
    Boston = BOSTON, MA;  Seattle = SEATTLE, WA;  Houston = HOUSTON, TX
    New York = NEW YORK / NEW YORK CITY / NYC / MANHATTAN, NY (other boroughs excluded)
    San Francisco Bay Area = the CA city list in BAY_AREA below
  Wage: WAGE_RATE_OF_PAY_FROM annualized (Hour x2080, Week x52, Bi-Weekly x26, Month x12, Year x1); median per group.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import statistics
import sys
import time
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "sponsorship" / "raw"
RAW_GLOB = "LCA_Disclosure_Data_FY2025_Q*.xlsx"
OUT = ROOT / "data" / "sponsorship" / "lca_filtered.csv"
PROV = ROOT / "data" / "sponsorship" / "provenance.json"
SOURCE_PAGE = "https://www.dol.gov/agencies/eta/foreign-labor/performance"
SOURCE_URL_PATTERN = "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/{name}"
FISCAL_YEAR = 2025
DOWNLOADED_ON = "2026-10-23"

BAY_AREA = {"SAN FRANCISCO", "SOUTH SAN FRANCISCO", "OAKLAND", "BERKELEY", "EMERYVILLE", "SAN JOSE", "SANTA CLARA", "SUNNYVALE",
            "MOUNTAIN VIEW", "PALO ALTO", "MENLO PARK", "REDWOOD CITY", "REDWOOD SHORES", "SAN MATEO", "FOSTER CITY", "BURLINGAME",
            "SAN BRUNO", "CUPERTINO", "MILPITAS", "FREMONT", "LOS ALTOS", "LOS GATOS", "CAMPBELL", "SAN CARLOS", "BELMONT",
            "PLEASANTON", "WALNUT CREEK", "SAN RAMON", "DUBLIN", "HAYWARD", "NEWARK", "UNION CITY", "SAN LEANDRO", "ALAMEDA"}
NYC = {"NEW YORK", "NEW YORK CITY", "NYC", "MANHATTAN"}
METRO_RULES = [("Boston", "MA", {"BOSTON"}), ("Seattle", "WA", {"SEATTLE"}), ("Houston", "TX", {"HOUSTON"}),
               ("New York", "NY", NYC), ("San Francisco Bay Area", "CA", BAY_AREA)]
PM_TITLE = re.compile(r"\bproduct\s+(manager|management|lead|owner|director)\b", re.I)
PM_EXCLUDE = re.compile(r"product\s+marketing", re.I)
ANNUAL = {"YEAR": 1, "HOUR": 2080, "WEEK": 52, "BI-WEEKLY": 26, "MONTH": 12}
COLUMNS = ["DECISION_DATE", "CASE_STATUS", "VISA_CLASS", "JOB_TITLE", "SOC_CODE", "SOC_TITLE", "EMPLOYER_NAME", "WORKSITE_CITY", "WORKSITE_STATE",
           "WAGE_RATE_OF_PAY_FROM", "WAGE_UNIT_OF_PAY", "TOTAL_WORKER_POSITIONS"]


def occupation(soc: str, title: str) -> str | None:
    soc = (soc or "").strip()
    if soc.startswith("13-1111"):
        return "management_consulting"
    if soc.startswith(("13-2051", "13-2054", "11-3031")) or (soc.startswith("13-1199") and re.search(r"strateg", title or "", re.I)):
        return "corporate_finance_strategy"
    if PM_TITLE.search(title or "") and not PM_EXCLUDE.search(title or "") and soc[:3] in ("11-", "13-", "15-"):
        return "product_management"
    return None


def metro(city: str, state: str) -> str | None:
    c, s = re.sub(r"\s+", " ", str(city or "").strip().upper()), str(state or "").strip().upper()
    for name, st, cities in METRO_RULES:
        if s == st and c in cities:
            return name
    return None


def annual_wage(v, unit) -> float | None:
    try:
        x = float(str(v).replace(",", "").replace("$", ""))
    except (TypeError, ValueError):
        return None
    m = ANNUAL.get(str(unit or "").strip().upper())
    return round(x * m, 2) if m and x > 0 else None


def keep(row: dict) -> tuple[str, str] | None:
    """Return (occupation, metro) if the case passes every filter, else None."""
    if str(row["CASE_STATUS"]).strip() != "Certified" or str(row["VISA_CLASS"]).strip() != "H-1B":
        return None
    occ, met = occupation(str(row["SOC_CODE"] or ""), str(row["JOB_TITLE"] or "")), metro(row["WORKSITE_CITY"], row["WORKSITE_STATE"])
    return (occ, met) if occ and met else None


def aggregate(rows) -> list[dict]:
    groups: dict[tuple, dict] = defaultdict(lambda: {"filings": 0, "positions": 0, "wages": []})
    for r in rows:
        k = keep(r)
        if not k:
            continue
        g = groups[(" ".join(str(r["EMPLOYER_NAME"]).split()), k[0], k[1])]
        g["filings"] += 1
        try:
            g["positions"] += int(r["TOTAL_WORKER_POSITIONS"] or 1)
        except (TypeError, ValueError):
            g["positions"] += 1
        w = annual_wage(r["WAGE_RATE_OF_PAY_FROM"], r["WAGE_UNIT_OF_PAY"])
        if w:
            g["wages"].append(w)
    out = []
    for (emp, occ, met), g in groups.items():
        out.append({"employer_name": emp, "occupation": occ, "metro": met, "fiscal_year": FISCAL_YEAR, "filings": g["filings"],
                    "worker_positions": g["positions"], "median_annual_wage": round(statistics.median(g["wages"])) if g["wages"] else ""})
    out.sort(key=lambda x: (x["metro"], x["occupation"], -x["filings"], x["employer_name"]))
    return out


def read_xlsx(path: Path):
    import openpyxl  # dev-only dependency (requirements-dev.txt); the app reads only the filtered CSV
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise SystemExit(f"Expected columns not found in {path.name}: {missing}")
    idx = {c: header.index(c) for c in COLUMNS}
    n = 0
    for vals in it:
        if all(v is None for v in vals):   # the sheet is padded with fully empty rows after the last case
            break
        n += 1
        yield {c: vals[i] for c, i in idx.items()}
    read_xlsx.count, read_xlsx.header = n, header


def main(min_filings: int = 1) -> None:
    files = sorted(RAW_DIR.glob(RAW_GLOB))
    if not files:
        raise SystemExit(f"No {RAW_GLOB} in {RAW_DIR}. Download from {SOURCE_PAGE} (see data/sponsorship/README.md).")
    t = time.time()
    status_counts: dict[str, int] = defaultdict(int)
    per_file = []

    def counted():
        for f in files:
            dmin = dmax = None
            for r in read_xlsx(f):
                status_counts[str(r["CASE_STATUS"])] += 1
                d = r["DECISION_DATE"]
                if d is not None:
                    dmin = d if dmin is None or d < dmin else dmin
                    dmax = d if dmax is None or d > dmax else dmax
                yield r
            per_file.append({"file_name": f.name, "source_url": SOURCE_URL_PATTERN.format(name=f.name),
                             "sha256": hashlib.sha256(f.read_bytes()).hexdigest(), "cases": read_xlsx.count,
                             "columns": len(read_xlsx.header),
                             "decision_dates": [dmin.date().isoformat() if dmin else None, dmax.date().isoformat() if dmax else None]})
    rows = [r for r in aggregate(counted()) if r["filings"] >= min_filings]
    quarters = sorted(re.search(r"_(Q\d)\.xlsx$", f.name).group(1) for f in files)
    coverage = ("full fiscal year (Q1–Q4)" if quarters == ["Q1", "Q2", "Q3", "Q4"]
                else f"PARTIAL fiscal year: {', '.join(quarters)} only")
    for r in rows:
        r["coverage"] = "FY2025" if quarters == ["Q1", "Q2", "Q3", "Q4"] else f"FY2025 {'+'.join(quarters)}"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)
    by = defaultdict(int)
    for r in rows:
        by[f"{r['metro']} | {r['occupation']}"] += r["filings"]
    PROV.write_text(json.dumps({
        "source": "U.S. Department of Labor, Office of Foreign Labor Certification — LCA Programs (H-1B, H-1B1, E-3) disclosure data",
        "source_page": SOURCE_PAGE, "fiscal_year": FISCAL_YEAR, "quarters_used": quarters, "coverage": coverage,
        "fiscal_year_note": "FY2025 = Oct 1, 2024 – Sep 30, 2025. Each quarterly file contains only that quarter's decisions.",
        "files": per_file, "downloaded_on": DOWNLOADED_ON,
        "downloaded_by": "the project owner, in a browser (dol.gov blocks automated downloads, HTTP 403)",
        "record_layout": "LCA_Record_Layout_FY2025_Q4.pdf",
        "case_status_counts": dict(status_counts), "columns_used": COLUMNS,
        "filters": {"visa_class": "H-1B", "case_status": "Certified", "metros": {m: sorted(c) for m, _, c in METRO_RULES},
                    "occupations": "see module docstring of scripts/build_sponsorship.py (stated assumptions)"},
        "output": {"file": "data/sponsorship/lca_filtered.csv", "rows": len(rows), "total_filings": sum(r["filings"] for r in rows),
                   "filings_by_metro_occupation": dict(sorted(by.items()))},
        "built_on": time.strftime("%Y-%m-%d"), "build_seconds": round(time.time() - t, 1),
        "caution": "Historical H-1B LCA filings by an employer. An LCA is not a visa approval and not a promise to sponsor any person or role.",
    }, indent=1), encoding="utf-8")
    print(f"{coverage}: {len(rows)} rows, {sum(r['filings'] for r in rows)} filings -> {OUT} ({OUT.stat().st_size / 1e3:.0f} KB) in {time.time() - t:.0f}s")


if __name__ == "__main__":
    main(int(sys.argv[1]) if len(sys.argv) > 1 else 1)
